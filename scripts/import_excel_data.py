import openpyxl,re,hashlib
from pathlib import Path

def q(v):
    return 'NULL' if v is None else "'" + str(v).replace("'","''") + "'"

def clean(v):
    return re.sub(r'\s+',' ',str(v or '')).strip()

def phone(v):
    s=str(v or '').strip().translate(str.maketrans('۰۱۲۳۴۵۶۷۸۹','0123456789'))
    s=re.sub(r'\D','',s)
    if len(s)==10 and s.startswith('9'):
        s='0'+s
    return s if re.fullmatch(r'09\d{9}',s) else None

ps=list(openpyxl.load_workbook('frontend/public/uploaded/products.xlsx',data_only=True).active.iter_rows(min_row=2,values_only=True))
rs=list(openpyxl.load_workbook('frontend/public/uploaded/orders.xlsx',data_only=True).active.iter_rows(min_row=2,values_only=True))

if len(ps)!=270:
    raise SystemExit(f'Unexpected product rows: {len(ps)}')

customers={}
for r in rs:
    p=phone(r[0])
    if p and p!='09153090907' and p not in customers:
        customers[p]=r

if len(customers)!=300:
    raise SystemExit(f'Unexpected valid customers after admin exclusion: {len(customers)}')

o=['PRAGMA foreign_keys=ON;']

o += [
"DELETE FROM payment_attempts WHERE payment_id IN (SELECT id FROM payments WHERE order_id IN (SELECT id FROM orders WHERE user_id='879eda82-3879-445e-b8a7-b75f3ecab43b'));",
"DELETE FROM payments WHERE order_id IN (SELECT id FROM orders WHERE user_id='879eda82-3879-445e-b8a7-b75f3ecab43b');",
"DELETE FROM coupon_usages WHERE order_id IN (SELECT id FROM orders WHERE user_id='879eda82-3879-445e-b8a7-b75f3ecab43b');",
"DELETE FROM order_notes WHERE order_id IN (SELECT id FROM orders WHERE user_id='879eda82-3879-445e-b8a7-b75f3ecab43b');",
"DELETE FROM shipping_tracking WHERE order_id IN (SELECT id FROM orders WHERE user_id='879eda82-3879-445e-b8a7-b75f3ecab43b');",
"DELETE FROM order_status_history WHERE order_id IN (SELECT id FROM orders WHERE user_id='879eda82-3879-445e-b8a7-b75f3ecab43b');",
"DELETE FROM stock_reservations WHERE order_id IN (SELECT id FROM orders WHERE user_id='879eda82-3879-445e-b8a7-b75f3ecab43b');",
"DELETE FROM order_items WHERE order_id IN (SELECT id FROM orders WHERE user_id='879eda82-3879-445e-b8a7-b75f3ecab43b');",
"DELETE FROM orders WHERE user_id='879eda82-3879-445e-b8a7-b75f3ecab43b';",
"DELETE FROM cart_items WHERE product_id IN ('b888d0aa-7ed2-4712-86bd-2c9a67ea5ec2','prod_demo','sample_mehr','sample_shab','sample_khak','sample_barg','sample_sokoot','sample_atiq','test_grid_01','test_grid_02','test_grid_03');",
"DELETE FROM favorites WHERE product_id IN ('b888d0aa-7ed2-4712-86bd-2c9a67ea5ec2','prod_demo','sample_mehr','sample_shab','sample_khak','sample_barg','sample_sokoot','sample_atiq','test_grid_01','test_grid_02','test_grid_03');",
"DELETE FROM reviews WHERE product_id IN ('b888d0aa-7ed2-4712-86bd-2c9a67ea5ec2','prod_demo','sample_mehr','sample_shab','sample_khak','sample_barg','sample_sokoot','sample_atiq','test_grid_01','test_grid_02','test_grid_03');",
"DELETE FROM product_attribute_assignments WHERE product_id IN ('b888d0aa-7ed2-4712-86bd-2c9a67ea5ec2','prod_demo','sample_mehr','sample_shab','sample_khak','sample_barg','sample_sokoot','sample_atiq','test_grid_01','test_grid_02','test_grid_03');",
"DELETE FROM discount_products WHERE product_id IN ('b888d0aa-7ed2-4712-86bd-2c9a67ea5ec2','prod_demo','sample_mehr','sample_shab','sample_khak','sample_barg','sample_sokoot','sample_atiq','test_grid_01','test_grid_02','test_grid_03');",
"DELETE FROM coupon_products WHERE product_id IN ('b888d0aa-7ed2-4712-86bd-2c9a67ea5ec2','prod_demo','sample_mehr','sample_shab','sample_khak','sample_barg','sample_sokoot','sample_atiq','test_grid_01','test_grid_02','test_grid_03');",
"DELETE FROM product_tag_map WHERE product_id IN ('b888d0aa-7ed2-4712-86bd-2c9a67ea5ec2','prod_demo','sample_mehr','sample_shab','sample_khak','sample_barg','sample_sokoot','sample_atiq','test_grid_01','test_grid_02','test_grid_03');",
"DELETE FROM related_products WHERE product_id IN ('b888d0aa-7ed2-4712-86bd-2c9a67ea5ec2','prod_demo','sample_mehr','sample_shab','sample_khak','sample_barg','sample_sokoot','sample_atiq','test_grid_01','test_grid_02','test_grid_03') OR related_product_id IN ('b888d0aa-7ed2-4712-86bd-2c9a67ea5ec2','prod_demo','sample_mehr','sample_shab','sample_khak','sample_barg','sample_sokoot','sample_atiq','test_grid_01','test_grid_02','test_grid_03');",
"DELETE FROM product_media WHERE product_id IN ('b888d0aa-7ed2-4712-86bd-2c9a67ea5ec2','prod_demo','sample_mehr','sample_shab','sample_khak','sample_barg','sample_sokoot','sample_atiq','test_grid_01','test_grid_02','test_grid_03');",
"DELETE FROM product_seo WHERE product_id IN ('b888d0aa-7ed2-4712-86bd-2c9a67ea5ec2','prod_demo','sample_mehr','sample_shab','sample_khak','sample_barg','sample_sokoot','sample_atiq','test_grid_01','test_grid_02','test_grid_03');",
"DELETE FROM product_variants WHERE product_id IN ('b888d0aa-7ed2-4712-86bd-2c9a67ea5ec2','prod_demo','sample_mehr','sample_shab','sample_khak','sample_barg','sample_sokoot','sample_atiq','test_grid_01','test_grid_02','test_grid_03');",
"DELETE FROM inventory_transactions WHERE product_id IN ('b888d0aa-7ed2-4712-86bd-2c9a67ea5ec2','prod_demo','sample_mehr','sample_shab','sample_khak','sample_barg','sample_sokoot','sample_atiq','test_grid_01','test_grid_02','test_grid_03');",
"DELETE FROM inventory WHERE product_id IN ('b888d0aa-7ed2-4712-86bd-2c9a67ea5ec2','prod_demo','sample_mehr','sample_shab','sample_khak','sample_barg','sample_sokoot','sample_atiq','test_grid_01','test_grid_02','test_grid_03');",
"DELETE FROM product_images WHERE product_id IN ('b888d0aa-7ed2-4712-86bd-2c9a67ea5ec2','prod_demo','sample_mehr','sample_shab','sample_khak','sample_barg','sample_sokoot','sample_atiq','test_grid_01','test_grid_02','test_grid_03');",
"DELETE FROM products WHERE id IN ('b888d0aa-7ed2-4712-86bd-2c9a67ea5ec2','prod_demo','sample_mehr','sample_shab','sample_khak','sample_barg','sample_sokoot','sample_atiq','test_grid_01','test_grid_02','test_grid_03');",
"DELETE FROM customer_group_members WHERE user_id IN ('usr_sample_01','usr_sample_02','usr_sample_03');",
"DELETE FROM customer_blocks WHERE user_id IN ('usr_sample_01','usr_sample_02','usr_sample_03');",
"DELETE FROM login_history WHERE user_id IN ('usr_sample_01','usr_sample_02','usr_sample_03');",
"DELETE FROM security_sessions WHERE user_id IN ('usr_sample_01','usr_sample_02','usr_sample_03');",
"DELETE FROM sessions WHERE user_id IN ('usr_sample_01','usr_sample_02','usr_sample_03');",
"DELETE FROM cart_items WHERE cart_id IN (SELECT id FROM carts WHERE user_id IN ('usr_sample_01','usr_sample_02','usr_sample_03'));",
"DELETE FROM carts WHERE user_id IN ('usr_sample_01','usr_sample_02','usr_sample_03');",
"DELETE FROM favorites WHERE user_id IN ('usr_sample_01','usr_sample_02','usr_sample_03');",
"DELETE FROM reviews WHERE user_id IN ('usr_sample_01','usr_sample_02','usr_sample_03');",
"DELETE FROM addresses WHERE user_id IN ('usr_sample_01','usr_sample_02','usr_sample_03');",
"DELETE FROM users WHERE id IN ('usr_sample_01','usr_sample_02','usr_sample_03');"
]

o += [
"DELETE FROM coupon_rules WHERE coupon_id IN ('coupon_sample_10','coupon_sample_fixed');",
"DELETE FROM coupons WHERE id IN ('coupon_sample_10','coupon_sample_fixed');"
]

all_stmts=list(o)

for n,r in enumerate(ps,1):
    name=clean(r[0]); desc=clean(r[1]); price=int(r[4]); pic=clean(r[5]); pid=f'xls_{n:04d}'
    all_stmts.append(
        f"INSERT OR IGNORE INTO products(id,category_id,slug,sku,name,description,price_irt,active) "
        f"VALUES({q(pid)},NULL,{q('product-'+str(n).zfill(4))},{q('GA-XLS-'+str(n).zfill(4))},{q(name)},{q(desc)},{price},1);"
    )
    all_stmts.append(
        f"INSERT OR IGNORE INTO product_images(id,product_id,path,alt_text,sort_order,is_primary) "
        f"VALUES({q('xlsimg_'+str(n).zfill(4))},{q(pid)},{q('/glsArt/uploaded/'+pic)},{q(name)},0,1);"
    )

for p,r in customers.items():
    uid='cust_'+hashlib.sha256(p.encode()).hexdigest()[:24]
    name=clean(str(r[1] or '')+' '+str(r[2] or ''))
    all_stmts.append(
        f"INSERT INTO users(id,mobile,name) VALUES({q(uid)},{q(p)},{q(name)}) "
        f"ON CONFLICT(mobile) DO UPDATE SET name=excluded.name,updated_at=CURRENT_TIMESTAMP;"
    )

parts=Path('import_parts')
parts.mkdir(exist_ok=True)
for old in parts.glob('*.sql'):
    old.unlink()

chunk=5
for i in range(0,len(all_stmts),chunk):
    batch=all_stmts[i:i+chunk]
    (parts/f'{i//chunk+1:04d}.sql').write_text('\n'.join(batch)+'\n',encoding='utf-8')

print(f'validated products={len(ps)} customers={len(customers)} parts={len(list(parts.glob("*.sql")))}')
